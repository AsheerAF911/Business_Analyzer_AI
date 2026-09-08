from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .base import BaseParser, ParsedRecord


def normalize_value(value: Any) -> Any:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.isoformat()

    if isinstance(value, date):
        return value.isoformat()

    return value


def clean_header(value: Any) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\n", " ")
        .split()
    ).strip()


class ExcelParser(BaseParser):

    def parse(
        self,
        file_path: Path,
        source_file: str,
    ) -> list[ParsedRecord]:

        workbook = load_workbook(
            file_path,
            data_only=True,
            read_only=True,
        )

        records: list[ParsedRecord] = []

        for worksheet in workbook.worksheets:

            headers = [
                clean_header(cell.value)
                for cell in worksheet[1]
            ]

            for row_number, values in enumerate(
                worksheet.iter_rows(
                    min_row=2,
                    values_only=True,
                ),
                start=2,
            ):

                # Ignore completely empty rows.
                if not any(
                    value not in (None, "")
                    for value in values
                ):
                    continue

                fields = {}

                for header, value in zip(
                    headers,
                    values,
                ):
                    if not header:
                        continue

                    fields[header] = normalize_value(
                        value
                    )

                records.append(
                    ParsedRecord(
                        source_file=source_file,
                        sheet=worksheet.title,
                        row_number=row_number,
                        fields=fields,
                    )
                )

        return records