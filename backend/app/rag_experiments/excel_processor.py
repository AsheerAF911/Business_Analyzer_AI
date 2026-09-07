from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPORT_TYPE = "inventory"


def normalize_value(value: Any) -> Any:
    """Convert Excel/Python values into JSON-friendly values."""
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def excel_date_to_iso(value: Any) -> str | None:
    """Return an ISO date when the cell contains a date-like value.

    openpyxl may expose an Excel serial as an int when the worksheet does not
    carry date formatting. The sample report stores DATE as serial 46143,
    which corresponds to 2026-05-01.
    """
    if value is None or value == "":
        return None

    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel's 1900 date system. openpyxl's helper avoids reimplementing
        # the leap-year compatibility behavior.
        from openpyxl.utils.datetime import from_excel

        converted = from_excel(value)
        if isinstance(converted, datetime):
            return converted.date().isoformat()
        if isinstance(converted, date):
            return converted.isoformat()

    return str(value)


def clean_header(header: Any) -> str:
    if header is None:
        return ""
    return " ".join(str(header).replace("\n", " ").split()).strip()


def row_to_record(headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
    record: dict[str, Any] = {}
    for header, value in zip(headers, values):
        if not header:
            continue
        record[header] = normalize_value(value)
    return record


def value(record: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in record and record[key] not in (None, ""):
            return record[key]
    return None


def business_text(record: dict[str, Any]) -> str:
    """Build business-readable text instead of dumping the raw row.

    The wording is intentionally explicit so a future retrieval/LLM step can
    understand what each value means without relying on column headers alone.
    """
    parts: list[str] = []

    date_value = excel_date_to_iso(value(record, "DATE"))
    inward = value(record, "INWARD NO")
    production = value(record, "PRODUCTION")
    vendor = value(record, "VENDOR NAME")
    product = value(record, "PRODUCT DESCRIPTION")
    barcode = value(record, "BARCODE")
    packaging = value(record, "PACKAGING SIZE (KG)")
    qty_pcs = value(record, "QUANTITY RECEIVED (NO. OF PCS.)")
    loose_qty = value(record, "LOOSE QTY")
    total_qty = value(record, "TOTAL QUANTITY RECEIVED (KG)")
    shortage = value(record, "SHORTAGE")
    net_stock = value(record, "NET STOCK QTY")
    inv_do = value(record, "INV/DO")
    invoice = value(record, "INVOICE")
    delivery = value(record, "DELIVERY NO")
    qcr = value(record, "QCR NO")
    batch = value(record, "BATCH NO")
    brand = value(record, "BRAND NAME")
    category = value(record, "CATEGORY (A, OR B OR C)")
    invoice_delivery_date = excel_date_to_iso(
        value(record, "INVOICE/DELIVERY DATE")
    )
    expiry_date = excel_date_to_iso(value(record, "EXPIRY DATE"))
    erp_receipt = value(record, "ERP RECIPET NO")

    if date_value:
        parts.append(f"Inventory transaction on {date_value}.")
    if inward:
        parts.append(f"Inward number: {inward}.")
    if production:
        parts.append(f"Production/location: {production}.")
    if vendor:
        parts.append(f"Vendor/source: {vendor}.")
    if product:
        parts.append(f"Product: {product}.")
    if barcode is not None:
        parts.append(f"Barcode: {barcode}.")
    if packaging is not None:
        parts.append(f"Packaging size: {packaging} kg.")
    if qty_pcs is not None:
        parts.append(f"Quantity received: {qty_pcs} pcs.")
    if loose_qty is not None:
        parts.append(f"Loose quantity: {loose_qty}.")
    if total_qty is not None:
        parts.append(f"Total quantity received: {total_qty} kg.")
    if shortage is not None:
        parts.append(f"Shortage: {shortage}.")
    if net_stock is not None:
        parts.append(f"Net stock quantity: {net_stock} kg.")
    if batch:
        parts.append(f"Batch number: {batch}.")
    if brand:
        parts.append(f"Brand: {brand}.")
    if inv_do:
        parts.append(f"Transaction document type: {inv_do}.")
    if invoice:
        parts.append(f"Invoice: {invoice}.")
    if delivery:
        parts.append(f"Delivery number: {delivery}.")
    if qcr:
        parts.append(f"QCR number: {qcr}.")
    if category:
        parts.append(f"Category: {category}.")
    if invoice_delivery_date:
        parts.append(f"Invoice/delivery date: {invoice_delivery_date}.")
    if expiry_date:
        parts.append(f"Expiry date: {expiry_date}.")
    if erp_receipt:
        parts.append(f"ERP receipt number: {erp_receipt}.")

    return " ".join(parts)


def make_metadata(
    record: dict[str, Any],
    *,
    source_file: str,
    sheet: str,
    row_number: int,
) -> dict[str, Any]:
    return {
        "source_file": source_file,
        "report_type": REPORT_TYPE,
        "sheet": sheet,
        "row_number": row_number,
        "date": excel_date_to_iso(value(record, "DATE")),
        "transaction_number": value(record, "INWARD NO"),
        "product": value(record, "PRODUCT DESCRIPTION"),
        "barcode": value(record, "BARCODE"),
        "vendor": value(record, "VENDOR NAME"),
    }


def process_sheet(
    worksheet,
    *,
    source_file: str,
) -> list[dict[str, Any]]:
    headers = [clean_header(cell.value) for cell in worksheet[1]]

    chunks: list[dict[str, Any]] = []

    # Business-aware strategy:
    # one meaningful transaction/item row is one atomic chunk. This keeps
    # product, quantity, inward number, vendor, and date together, while
    # preserving precise row-level evidence. It is easier to change this
    # function later if domain-specific grouping proves useful.
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        if not any(value not in (None, "") for value in values):
            continue

        record = row_to_record(headers, values)
        text = business_text(record)
        if not text:
            continue

        chunks.append(
            {
                "chunk_id": f"{Path(source_file).stem}-{worksheet.title}-{row_number}",
                "text": text,
                "metadata": make_metadata(
                    record,
                    source_file=Path(source_file).name,
                    sheet=worksheet.title,
                    row_number=row_number,
                ),
            }
        )

    return chunks


def process_workbook(input_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(input_path, data_only=True, read_only=True)

    chunks: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        chunks.extend(
            process_sheet(
                worksheet,
                source_file=input_path.name,
            )
        )

    return chunks


def write_json(chunks: list[dict[str, Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as file:
        json.dump(chunks, file, indent=2, ensure_ascii=False)


def inspect_chunks(chunks: list[dict[str, Any]], preview_count: int = 5) -> None:
    print(f"Total records: {len(chunks)}")
    print(f"Total chunks: {len(chunks)}")
    print()

    for index, chunk in enumerate(chunks[:preview_count], start=1):
        print(f"--- Chunk {index} ---")
        print(f"chunk_id: {chunk['chunk_id']}")
        print(f"text: {chunk['text']}")
        print("metadata:")
        print(json.dumps(chunk["metadata"], indent=2, ensure_ascii=False))
        print()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Educational Day 4 inventory Excel-to-RAG-chunks experiment."
    )
    parser.add_argument("input", type=Path, help="Path to the Excel report.")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("inventory_chunks.json"),
        help="Path for the generated JSON file.",
    )
    parser.add_argument(
        "--preview",
        type=int,
        default=5,
        help="Number of chunks to print for manual inspection.",
    )
    args = parser.parse_args()

    chunks = process_workbook(args.input)
    write_json(chunks, args.output)
    inspect_chunks(chunks, args.preview)

    print(f"JSON written to: {args.output}")


if __name__ == "__main__":
    main()
